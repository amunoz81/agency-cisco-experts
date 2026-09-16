"""Exportación de la propuesta a PowerPoint (deck ejecutivo) y Word (documento
técnico) con la identidad Cisco: fondo navy oscuro, texto blanco, acentos verde
y cian, logo Cisco y pie de copyright. Requiere el extra `office`/`web`
(python-pptx, python-docx, Pillow). Best-effort: si falta una librería, lanza
ImportError que el llamador captura.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from ..schemas import (
    Bom,
    CritiqueResult,
    ExecutiveSynthesis,
    FinancialCase,
    Opportunity,
    ReviewResult,
    ScopeClassification,
    SpecialistFinding,
)

ARCH_LABELS = {
    "secure_networking": "Secure Networking",
    "security": "Security",
    "it_ot": "IT/OT (Cyber Vision)",
    "observability_soc": "Observabilidad y SOC",
    "datacenter_ai": "Data Center y AI",
    "collaboration": "Collaboration (Webex)",
}

# Paleta Cisco (según la identidad de las láminas)
DECK_BG = (0x0A, 0x1A, 0x2E)   # navy muy oscuro (fondo)
CARD_BG = (0x12, 0x24, 0x3C)   # navy un poco más claro (tarjetas/filas)
ROW_ALT = (0x0E, 0x1E, 0x34)
GREEN = (0x1A, 0xB2, 0x4B)     # verde Cisco
CYAN = (0x00, 0xBC, 0xEB)      # cian Cisco
WHITE = (0xFF, 0xFF, 0xFF)
LIGHT = (0xC7, 0xD6, 0xE5)     # texto atenuado sobre navy
PRIMARY = (0x0D, 0x27, 0x4D)
INK = (0x1A, 0x1A, 0x1A)
MUTED = (0x5B, 0x6B, 0x7A)

COPYRIGHT = f"© {datetime.now().year} Cisco and/or its affiliates. All rights reserved."


def _slug(text: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in (text or "cliente").lower()).strip("_")


def _cls_label(v) -> str:
    v = v.value if isinstance(v, ScopeClassification) else v
    return {
        "necesaria": "Necesaria",
        "opcional_justificada": "Opcional justificada",
        "fuera_de_alcance": "Fuera de alcance",
    }.get(v, str(v))


def _star(opportunity: Opportunity, synthesis: ExecutiveSynthesis | None) -> dict:
    if synthesis:
        return {
            "situation": synthesis.situation,
            "task": synthesis.task,
            "action": synthesis.action,
            "result": synthesis.result,
            "summary": synthesis.executive_summary,
            "contradictions": synthesis.contradictions_resolved,
        }
    objectives = "; ".join(opportunity.objectives) or "objetivos por confirmar"
    return {
        "situation": f"{opportunity.customer} requiere modernizar su arquitectura "
        "de red y seguridad de extremo a extremo.",
        "task": f"Objetivos: {objectives}.",
        "action": "Integración de las arquitecturas Cisco necesarias en un diseño único.",
        "result": "Arquitectura Zero Trust con BOM y caso financiero, lista para validación.",
        "summary": "",
        "contradictions": [],
    }


def _bars_png(out_dir: Path, color=WHITE) -> str:
    """Genera (una vez) el PNG de las barras del logo Cisco, sin texto (para no
    depender de fuentes del sistema). Devuelve la ruta."""
    from PIL import Image, ImageDraw

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "_cisco_bars.png"
    if path.exists():
        return str(path)
    heights = [20, 36, 60, 36, 20, 36, 60, 36, 20]
    bw, gap, x0, base = 10, 9, 6, 66
    w = x0 * 2 + len(heights) * bw + (len(heights) - 1) * gap
    img = Image.new("RGBA", (w, base + 10), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    x = x0
    for h in heights:
        d.rounded_rectangle([x, base - h, x + bw, base], radius=4, fill=(*color, 255))
        x += bw + gap
    img.save(path)
    return str(path)


# ==========================================================================
# PowerPoint (deck ejecutivo, tema Cisco oscuro)
# ==========================================================================
def build_pptx(
    *,
    opportunity: Opportunity,
    scope_plan: dict[str, ScopeClassification],
    findings: list[SpecialistFinding],
    integration: dict,
    bom: Bom,
    financial_case: FinancialCase,
    review: ReviewResult,
    verification: dict,
    fiscal_year: str,
    out_dir: str | Path,
    basename: str | None = None,
    synthesis: ExecutiveSynthesis | None = None,
) -> str:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    out = Path(out_dir)
    bars_png = _bars_png(out)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SW, SH = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def rgb(c):
        return RGBColor(*c)

    def box(slide, shape, x, y, w, h, fill=None, line=None, line_w=1.0):
        shp = slide.shapes.add_shape(shape, x, y, w, h)
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
        shp.shadow.inherit = False
        return shp

    def tframe(slide, x, y, w, h):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tb.text_frame.word_wrap = True
        return tb.text_frame

    def para(tf, text, size=14, color=WHITE, bold=False, first=False, bullet=False, space=2):
        p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
        p.space_after = Pt(space)
        run = p.add_run()
        run.text = ("•  " + text) if bullet else text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        return p

    def background(slide):
        box(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=DECK_BG)
        # marco tenue
        box(slide, MSO_SHAPE.RECTANGLE, Inches(0.12), Inches(0.12),
            SW - Inches(0.24), SH - Inches(0.24), fill=None, line=(0x2A, 0x3C, 0x52), line_w=0.75)

    def logo_and_footer(slide):
        # Logo Cisco (barras + CISCO) abajo a la derecha
        slide.shapes.add_picture(bars_png, SW - Inches(1.15), SH - Inches(0.78),
                                 height=Inches(0.30))
        tf = tframe(slide, SW - Inches(1.25), SH - Inches(0.5), Inches(1.1), Inches(0.3))
        para(tf, "C I S C O", size=10, color=WHITE, bold=True, first=True)
        # Copyright abajo a la izquierda
        tf = tframe(slide, Inches(0.35), SH - Inches(0.5), Inches(8), Inches(0.3))
        para(tf, COPYRIGHT, size=8, color=(0x7A, 0x8B, 0xA0), first=True)

    def header(slide, title, accent=CYAN):
        tf = tframe(slide, Inches(0.55), Inches(0.35), SW - Inches(1.1), Inches(0.8))
        para(tf, title, size=27, color=WHITE, bold=True, first=True)
        box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.15),
            Inches(1.4), Pt(5), fill=accent)

    def card(slide, x, y, w, h, accent=CYAN):
        return box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
                   fill=CARD_BG, line=accent, line_w=1.25)

    def content_slide(title, accent=CYAN):
        s = prs.slides.add_slide(blank)
        background(s)
        header(s, title, accent)
        logo_and_footer(s)
        return s

    def add_table(slide, x, y, w, headers, rows, accent=CYAN):
        tbl = slide.shapes.add_table(len(rows) + 1, len(headers), x, y, w, Inches(0.4)).table
        for j, htext in enumerate(headers):
            c = tbl.cell(0, j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(accent)
            r = c.text_frame.paragraphs[0].add_run()
            r.text = htext
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = rgb(DECK_BG)
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                c = tbl.cell(i, j)
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(CARD_BG if i % 2 else ROW_ALT)
                r = c.text_frame.paragraphs[0].add_run()
                r.text = str(val)
                r.font.size = Pt(10)
                r.font.color.rgb = rgb(WHITE)
        return tbl

    star = _star(opportunity, synthesis)

    # --- 1. Portada ---
    s = prs.slides.add_slide(blank)
    background(s)
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.55), Inches(1.7), Pt(7), fill=GREEN)
    tf = tframe(s, Inches(0.8), Inches(1.5), SW - Inches(1.6), Inches(0.8))
    para(tf, "PROPUESTA TÉCNICA Y COMERCIAL", size=16, color=CYAN, bold=True, first=True)
    tf = tframe(s, Inches(0.8), Inches(2.8), SW - Inches(1.6), Inches(1.6))
    para(tf, opportunity.customer, size=46, color=WHITE, bold=True, first=True)
    tf = tframe(s, Inches(0.8), Inches(4.5), SW - Inches(1.6), Inches(1.2))
    sub = f"Arquitectura Cisco de extremo a extremo · {fiscal_year}"
    if opportunity.industry:
        sub += f" · {opportunity.industry}"
    para(tf, sub, size=16, color=LIGHT, first=True)
    para(tf, "Agencia de Expertos Cisco · " + datetime.now().strftime("%Y-%m-%d"),
         size=12, color=(0x7A, 0x8B, 0xA0))
    logo_and_footer(s)

    # --- 2. Resumen ejecutivo ---
    s = content_slide("Resumen ejecutivo", GREEN)
    card(s, Inches(0.5), Inches(1.45), SW - Inches(1), Inches(5.15), GREEN)
    tf = tframe(s, Inches(0.8), Inches(1.65), SW - Inches(1.6), Inches(4.8))
    if star["summary"]:
        para(tf, star["summary"], size=14, color=CYAN, bold=True, first=True)
    para(tf, "Situación: " + star["situation"], size=13, bullet=True, first=not star["summary"])
    para(tf, "Tarea: " + star["task"], size=13, bullet=True)
    para(tf, "Acción: " + star["action"], size=13, bullet=True)
    para(tf, "Resultado: " + star["result"], size=13, bullet=True)
    for c in star["contradictions"][:3]:
        para(tf, "Contradicción resuelta: " + c, size=12, color=LIGHT, bullet=True)

    # --- 3. Alcance ---
    s = content_slide("Alcance de la solución", CYAN)
    rows = [[ARCH_LABELS.get(a, a), _cls_label(c)] for a, c in scope_plan.items()]
    add_table(s, Inches(0.6), Inches(1.5), Inches(8.6), ["Arquitectura", "Clasificación"], rows)

    # --- 4. Arquitecturas ---
    s = content_slide("Arquitecturas de la solución", GREEN)
    card(s, Inches(0.5), Inches(1.45), SW - Inches(1), Inches(5.15), GREEN)
    tf = tframe(s, Inches(0.8), Inches(1.6), SW - Inches(1.6), Inches(4.9))
    first = True
    for f in findings:
        para(tf, ARCH_LABELS.get(f.architecture.value, f.architecture.value),
             size=14, color=CYAN, bold=True, first=first, space=0)
        first = False
        para(tf, f.customer_need, size=10, color=LIGHT, space=0)
        para(tf, "→ " + f.proposed_solution[:150], size=10, space=6)

    # --- 5. Integración ---
    s = content_slide("Integración entre arquitecturas", CYAN)
    card(s, Inches(0.5), Inches(1.45), SW - Inches(1), Inches(5.15), CYAN)
    tf = tframe(s, Inches(0.8), Inches(1.65), SW - Inches(1.6), Inches(4.8))
    items = (integration.get("data_flows", []) + integration.get("policy_enforcement_points", []))
    if not items:
        items = ["Diseño integrado con identidad y política consistentes."]
    for i, x in enumerate(items):
        para(tf, x, size=12, bullet=True, first=(i == 0))

    # --- 6. BOM ---
    s = content_slide("BOM y licenciamiento (resumen)", GREEN)
    hw = sum(1 for line in bom.lines if not line.is_license)
    lic = sum(1 for line in bom.lines if line.is_license)
    tf = tframe(s, Inches(0.55), Inches(1.3), SW - Inches(1.1), Inches(0.4))
    para(tf, f"{len(bom.lines)} líneas · {hw} hardware · {lic} licencias/suscripciones",
         size=13, color=CYAN, bold=True, first=True)
    rows = [
        [ln.sku or "—", ln.description[:44], int(ln.quantity),
         ARCH_LABELS.get(ln.architecture.value, ln.architecture.value)]
        for ln in bom.lines[:11]
    ]
    add_table(s, Inches(0.55), Inches(1.85), Inches(12.2),
              ["SKU", "Descripción", "Cant.", "Arquitectura"], rows, GREEN)

    # --- 7. Caso financiero ---
    s = content_slide("Caso financiero", CYAN)
    scen = {sc.name: sc for sc in financial_case.scenarios}
    rows = []
    for name, r in financial_case.results.items():
        sc = scen.get(name)
        rows.append([
            name,
            f"{sc.capex:,.0f}" if sc else "-",
            f"{r.get('tco', 0):,.0f}",
            f"{r.get('roi', 0):.0%}",
            f"{r.get('npv', 0):,.0f}",
            (f"{r.get('payback_years')} años"
             if r.get("payback_years", -1) >= 0 else "> horizonte"),
        ])
    add_table(s, Inches(0.6), Inches(1.5), Inches(9.6),
              ["Escenario", "CAPEX", "TCO", "ROI", "NPV", "Payback"], rows)
    tf = tframe(s, Inches(0.6), Inches(3.7), SW - Inches(1.2), Inches(2.9))
    para(tf, f"Moneda: {financial_case.currency}. Supuestos:", size=12, color=LIGHT, first=True)
    for a in financial_case.assumptions:
        para(tf, a, size=10, color=LIGHT, bullet=True)

    # --- 8. Revisión y próximos pasos ---
    s = content_slide("Revisión técnica y próximos pasos", GREEN)
    card(s, Inches(0.5), Inches(1.45), SW - Inches(1), Inches(5.15), GREEN)
    tf = tframe(s, Inches(0.8), Inches(1.65), SW - Inches(1.6), Inches(4.8))
    estado = "APROBADA (sin bloqueantes)" if review.passed else "CON BLOQUEANTES"
    para(tf, "Estado de la revisión: " + estado, size=14, color=CYAN, bold=True, first=True)
    para(tf, f"Cobertura de evidencia verificada: {verification.get('coverage_ratio', 0):.0%}",
         size=12, color=LIGHT)
    if opportunity.pending_data:
        para(tf, "Datos pendientes del cliente:", size=13, color=GREEN, bold=True)
        for d in opportunity.pending_data:
            para(tf, d, size=11, bullet=True)

    out.mkdir(parents=True, exist_ok=True)
    path = out / f"propuesta_{_slug(basename or opportunity.customer)}.pptx"
    prs.save(str(path))
    return str(path)


# ==========================================================================
# Word (documento técnico con banner Cisco)
# ==========================================================================
def build_docx(
    *,
    opportunity: Opportunity,
    scope_plan: dict[str, ScopeClassification],
    findings: list[SpecialistFinding],
    integration: dict,
    bom: Bom,
    financial_case: FinancialCase,
    review: ReviewResult,
    verification: dict,
    fiscal_year: str,
    out_dir: str | Path,
    basename: str | None = None,
    synthesis: ExecutiveSynthesis | None = None,
    critique: CritiqueResult | None = None,
) -> str:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches as DInches
    from docx.shared import Pt, RGBColor

    out = Path(out_dir)
    bars_png = _bars_png(out)

    doc = Document()
    navy = RGBColor(*PRIMARY)

    def _hex(c):
        return "{:02X}{:02X}{:02X}".format(*c)

    def shade(cell, color):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), _hex(color))
        cell._tc.get_or_add_tcPr().append(shd)

    def h(text, level=1):
        p = doc.add_heading(text, level=level)
        for run in p.runs:
            run.font.color.rgb = navy
        return p

    def kv(label, value):
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(str(value))
        return p

    star = _star(opportunity, synthesis)

    # --- Banner Cisco (navy con logo + título) ---
    banner = doc.add_table(rows=1, cols=1)
    cell = banner.rows[0].cells[0]
    shade(cell, DECK_BG)
    p = cell.paragraphs[0]
    p.add_run().add_picture(bars_png, height=DInches(0.34))
    rc = p.add_run("  CISCO")
    rc.bold = True
    rc.font.size = Pt(16)
    rc.font.color.rgb = RGBColor(*WHITE)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(opportunity.customer)
    r2.bold = True
    r2.font.size = Pt(26)
    r2.font.color.rgb = RGBColor(*WHITE)
    p3 = cell.add_paragraph()
    r3 = p3.add_run(
        f"Propuesta técnica y comercial · Arquitectura Cisco de extremo a extremo · {fiscal_year}"
    )
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(*CYAN)

    doc.add_paragraph()
    if opportunity.industry:
        kv("Industria", opportunity.industry)
    doc.add_paragraph(datetime.now().strftime("%Y-%m-%d")).runs[0].font.color.rgb = RGBColor(*MUTED)

    # 1. Resumen ejecutivo
    h("1. Resumen ejecutivo (STAR)", 1)
    if star["summary"]:
        doc.add_paragraph(star["summary"])
    kv("Situación", star["situation"])
    kv("Tarea", star["task"])
    kv("Acción", star["action"])
    kv("Resultado", star["result"])
    if star["contradictions"]:
        h("Contradicciones resueltas por el coordinador", 2)
        for c in star["contradictions"]:
            doc.add_paragraph(c, style="List Bullet")

    def styled_table(headers):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for j, htext in enumerate(headers):
            shade(t.rows[0].cells[j], PRIMARY)
            run = t.rows[0].cells[j].paragraphs[0].add_run(htext)
            run.bold = True
            run.font.color.rgb = RGBColor(*WHITE)
        return t

    # 2. Alcance
    h("2. Alcance de la solución", 1)
    t = styled_table(["Arquitectura", "Clasificación"])
    for a, c in scope_plan.items():
        cells = t.add_row().cells
        cells[0].text = ARCH_LABELS.get(a, a)
        cells[1].text = _cls_label(c)

    # 3. Blueprint por arquitectura
    h("3. Blueprint por arquitectura", 1)
    for f in findings:
        h(ARCH_LABELS.get(f.architecture.value, f.architecture.value)
          + f"  [{_cls_label(f.scope)}]", 2)
        kv("Necesidad", f.customer_need)
        kv("Solución propuesta", f.proposed_solution)
        if f.dependencies:
            kv("Dependencias", "; ".join(f.dependencies))
        kv("Dimensionamiento", f.sizing)
        if f.licenses:
            kv("Licencias", "; ".join(f.licenses))
        kv("Beneficio medible", f.measurable_benefit)
        if f.risks_pending:
            kv("Riesgos y pendientes", "; ".join(f.risks_pending))
        if f.evidence:
            doc.add_paragraph().add_run("Evidencia:").bold = True
            for e in f.evidence:
                line = f"{e.claim} — {e.source}"
                if e.version:
                    line += f" ({e.version})"
                line += f" [{e.status.value}]"
                doc.add_paragraph(line, style="List Bullet")

    # 4. Integración
    h("4. Integración entre arquitecturas", 1)
    for key, label in (
        ("data_flows", "Flujos de datos"),
        ("policy_enforcement_points", "Puntos de aplicación de política"),
        ("operational_responsibilities", "Responsabilidades operativas"),
    ):
        items = integration.get(key, [])
        if items:
            h(label, 2)
            for x in items:
                doc.add_paragraph(x, style="List Bullet")

    # 5. BOM
    h("5. BOM y licenciamiento (consolidado)", 1)
    t = styled_table(["SKU", "Descripción", "Cant.", "Arquitectura", "Tipo", "Gestión"])
    for ln in bom.lines:
        cells = t.add_row().cells
        cells[0].text = ln.sku or "—"
        cells[1].text = ln.description
        cells[2].text = f"{int(ln.quantity)} {ln.unit}"
        cells[3].text = ARCH_LABELS.get(ln.architecture.value, ln.architecture.value)
        cells[4].text = "Licencia" if ln.is_license else "Hardware"
        cells[5].text = ln.management or "—"

    # 6. Caso financiero
    h("6. Caso financiero", 1)
    scen = {sc.name: sc for sc in financial_case.scenarios}
    t = styled_table(["Escenario", "CAPEX", "TCO", "ROI", "NPV", "Payback"])
    for name, r in financial_case.results.items():
        sc = scen.get(name)
        cells = t.add_row().cells
        cells[0].text = name
        cells[1].text = f"{sc.capex:,.0f}" if sc else "-"
        cells[2].text = f"{r.get('tco', 0):,.0f}"
        cells[3].text = f"{r.get('roi', 0):.0%}"
        cells[4].text = f"{r.get('npv', 0):,.0f}"
        cells[5].text = (
            f"{r.get('payback_years')} años" if r.get("payback_years", -1) >= 0 else "> horizonte"
        )
    kv("Moneda", financial_case.currency)
    doc.add_paragraph().add_run("Supuestos:").bold = True
    for a in financial_case.assumptions:
        doc.add_paragraph(a, style="List Bullet")

    # 7. Revisión técnica
    h("7. Revisión técnica independiente", 1)
    kv("Estado", "Aprobada (sin bloqueantes)" if review.passed else "Con bloqueantes")
    if review.issues:
        t = styled_table(["Severidad", "Categoría", "Detalle"])
        for i in review.issues:
            cells = t.add_row().cells
            cells[0].text = i.severity
            cells[1].text = i.category
            cells[2].text = i.detail
    if critique and critique.rationale:
        h("Crítica cualitativa (revisor)", 2)
        doc.add_paragraph(critique.rationale)
        for i in critique.issues:
            doc.add_paragraph(i, style="List Bullet")

    # 8. Métricas y pendientes
    h("8. Métricas de éxito y pendientes", 1)
    kv("Evidencia", f"verificada {verification.get('verified', 0)} · "
       f"condicionada {verification.get('conditioned', 0)} · "
       f"pendiente {verification.get('pending', 0)} "
       f"(cobertura {verification.get('coverage_ratio', 0):.0%})")
    if opportunity.pending_data:
        h("Datos pendientes del cliente", 2)
        for d in opportunity.pending_data:
            doc.add_paragraph(d, style="List Bullet")

    doc.add_paragraph()
    foot = doc.add_paragraph(COPYRIGHT + "  Documento generado por la Agencia de Expertos "
                             "Cisco; requiere revisión de un arquitecto humano.")
    foot.runs[0].font.size = Pt(9)
    foot.runs[0].font.color.rgb = RGBColor(*MUTED)

    out.mkdir(parents=True, exist_ok=True)
    path = out / f"propuesta_{_slug(basename or opportunity.customer)}.docx"
    doc.save(str(path))
    return str(path)
