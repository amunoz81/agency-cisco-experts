"""Lectura de la base instalada subida por el usuario (Excel / PDF / CSV).

Extrae líneas legibles que los agentes leen como inventario. No falla la
solicitud si una dependencia no está: devuelve lo que pueda y una nota.
"""

from __future__ import annotations

import csv
import io

MAX_LINES = 500


def parse_install_base(filename: str, content: bytes) -> dict:
    """Devuelve {items: [str], text: str, note: str} a partir del archivo."""
    name = (filename or "").lower()
    try:
        if name.endswith((".xlsx", ".xlsm")):
            return _parse_xlsx(content)
        if name.endswith(".csv"):
            return _parse_csv(content)
        if name.endswith(".pdf"):
            return _parse_pdf(content)
        if name.endswith((".txt", ".tsv")):
            text = content.decode("utf-8", errors="ignore")
            items = [ln.strip() for ln in text.splitlines() if ln.strip()][:MAX_LINES]
            return {"items": items, "text": text[:20000], "note": ""}
    except Exception as exc:  # pragma: no cover - depende del archivo
        return {"items": [], "text": "", "note": f"No se pudo leer el archivo: {exc}"}
    return {
        "items": [],
        "text": "",
        "note": f"Formato no soportado: {filename}. Usa .xlsx, .csv o .pdf.",
    }


def _parse_xlsx(content: bytes) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"items": [], "text": "", "note": "Instala el extra web (openpyxl) para leer Excel."}
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    items: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                items.append(" | ".join(cells))
            if len(items) >= MAX_LINES:
                break
        if len(items) >= MAX_LINES:
            break
    return {"items": items, "text": "\n".join(items)[:20000], "note": ""}


def _parse_csv(content: bytes) -> dict:
    text = content.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    items = []
    for row in reader:
        cells = [c.strip() for c in row if c and c.strip()]
        if cells:
            items.append(" | ".join(cells))
        if len(items) >= MAX_LINES:
            break
    return {"items": items, "text": "\n".join(items)[:20000], "note": ""}


def _parse_pdf(content: bytes) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError:
        return {
            "items": [], "text": "",
            "note": "Instala el extra web/ingest (pypdf) para leer PDF.",
        }
    reader = PdfReader(io.BytesIO(content))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    items = [ln.strip() for ln in text.splitlines() if ln.strip()][:MAX_LINES]
    return {"items": items, "text": text[:20000], "note": ""}
